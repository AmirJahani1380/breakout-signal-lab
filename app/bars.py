from __future__ import annotations

import csv
import hashlib
import os
import sqlite3
import tempfile
import time
from bisect import bisect_left
from collections.abc import Iterable, Sequence
from contextlib import closing
from dataclasses import dataclass
from datetime import UTC, datetime
from itertools import chain
from math import isfinite
from pathlib import Path
from typing import Any, NoReturn
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook

REQUIRED_PRICE_COLUMNS = ("time", "open", "high", "low", "close")
SUPPORTED_SOURCE_SUFFIXES = frozenset({".csv", ".parquet", ".xlsx"})
EXPORT_NAME_SUFFIXES = ("_max_bars", "-max-bars")
PARQUET_CACHE_VERSION = 1
PARQUET_CACHE_CLEANUP_LIMIT = 256


class SourceValidationError(ValueError):
    """A source workbook cannot safely be used."""


class _DecodedBinaryLines:
    """Expose decoded physical lines while retaining byte positions for csv.reader."""

    def __init__(self, source_file: Any, encoding: str = "utf-8") -> None:
        self.source_file = source_file
        self.encoding = encoding

    def __iter__(self) -> _DecodedBinaryLines:
        return self

    def __next__(self) -> str:
        line = self.source_file.readline()
        if not line:
            raise StopIteration
        return str(line.decode(self.encoding))

    @property
    def position(self) -> int:
        return int(self.source_file.tell())


@dataclass(frozen=True, slots=True)
class Bar:
    time: int
    open: float
    high: float
    low: float
    close: float
    volume: float

    def as_dict(self) -> dict[str, int | float]:
        return {
            "time": self.time,
            "open": self.open,
            "high": self.high,
            "low": self.low,
            "close": self.close,
            "volume": self.volume,
        }


class BarStore:
    def __init__(
        self,
        bars: Iterable[Bar],
        imported_features: pd.DataFrame | None = None,
        dataset_id: str = "",
    ) -> None:
        self.bars = tuple(bars)
        self.times = tuple(bar.time for bar in self.bars)
        self.dataset_id = dataset_id
        self.imported_features = _validated_imported_features(
            imported_features, self.times, dataset_id
        )

    @property
    def canonical_frame(self) -> pd.DataFrame:
        return pd.DataFrame(
            [bar.as_dict() for bar in self.bars],
            columns=("time", "open", "high", "low", "close", "volume"),
        ).set_index("time")

    def page(self, before: int | None, limit: int) -> tuple[tuple[Bar, ...], bool]:
        end = len(self.bars) if before is None else bisect_left(self.times, before)
        start = max(0, end - limit)
        return self.bars[start:end], start > 0


@dataclass(frozen=True, slots=True)
class BarPage:
    """A bounded calculation window and its display slice."""

    bars: tuple[Bar, ...]
    display_start: int
    has_more: bool

    @property
    def display_bars(self) -> tuple[Bar, ...]:
        return self.bars[self.display_start :]


class PagedBarReader:
    """Read bounded chronological windows without materializing a complete dataset."""

    def __init__(self, source_path: Path, source_timezone: str = "UTC") -> None:
        self.source_path = source_path
        self.source_timezone = _source_timezone(source_path, source_timezone)
        self._timestamps: tuple[int, ...] | None = None
        self._csv_offsets: tuple[int, ...] | None = None
        self._parquet_cache_path: Path | None = None
        if source_path.suffix.lower() == ".parquet" and source_path.is_file():
            self._parquet_cache_path = self._prepare_parquet_cache()

    def page(self, before: int | None, display_limit: int, warm_up: int) -> BarPage:
        if not self.source_path.is_file():
            raise SourceValidationError(f"{self.source_path.name}: source file does not exist")
        suffix = self.source_path.suffix.lower()
        if suffix == ".csv":
            return self._csv_page(before, display_limit, warm_up)
        if suffix == ".parquet":
            return self._parquet_page(before, display_limit, warm_up)
        if suffix == ".xlsx":
            return self._xlsx_page(before, display_limit, warm_up)
        raise SourceValidationError(f"{self.source_path.name}: unsupported source format")

    def _csv_page(self, before: int | None, display_limit: int, warm_up: int) -> BarPage:
        timestamps = self._timestamp_index()
        offsets = self._csv_offsets
        assert offsets is not None
        end = len(timestamps) if before is None else bisect_left(timestamps, before)
        return self._page_from_rows(
            self._read_csv_rows(offsets, max(0, end - display_limit - warm_up), end),
            end,
            display_limit,
            warm_up,
        )

    def _build_csv_index(self) -> tuple[int, ...]:
        try:
            with self.source_path.open("rb") as source_file:
                lines = _DecodedBinaryLines(source_file, "utf-8-sig")
                reader = csv.reader(lines)
                header = tuple(next(reader, []))
                normalized_header = tuple(name.strip() for name in header)
                if len(normalized_header) != len(set(normalized_header)):
                    raise SourceValidationError(
                        f"{self.source_path.name}: duplicate column names are not allowed"
                    )
                columns = _columns(self.source_path, normalized_header)
                offsets: list[int] = []
                timestamps: list[int] = []
                while True:
                    record_start = lines.position
                    try:
                        row = tuple(next(reader))
                    except StopIteration:
                        break
                    record_number = len(timestamps) + 2
                    if len(row) != len(header):
                        raise SourceValidationError(
                            f"{self.source_path.name}, CSV record {record_number}: expected "
                            f"{len(header)} fields, found {len(row)}"
                        )
                    offsets.append(record_start)
                    timestamps.append(
                        _timestamp(
                            self.source_path,
                            record_number,
                            _cell(row, columns["time"]),
                            self.source_timezone,
                        )
                    )
        except (OSError, UnicodeError, csv.Error, StopIteration) as error:
            raise SourceValidationError(
                f"{self.source_path.name}: cannot index CSV: {error}"
            ) from error
        self._csv_offsets = tuple(offsets)
        return tuple(timestamps)

    def _read_csv_rows(
        self, offsets: tuple[int, ...], start: int, end: int
    ) -> list[tuple[Any, ...]]:
        try:
            with self.source_path.open("rb") as source_file:
                header = next(csv.reader(_DecodedBinaryLines(source_file, "utf-8-sig")), [])
                rows = []
                for row_index in range(start, end):
                    source_file.seek(offsets[row_index])
                    rows.append(tuple(next(csv.reader(_DecodedBinaryLines(source_file)))))
        except (OSError, UnicodeError, csv.Error, StopIteration) as error:
            raise SourceValidationError(
                f"{self.source_path.name}: cannot read CSV page: {error}"
            ) from error
        return [tuple(header), *rows]

    def _parquet_page(self, before: int | None, display_limit: int, warm_up: int) -> BarPage:
        timestamps = self._timestamp_index()
        end = len(timestamps) if before is None else bisect_left(timestamps, before)
        start = max(0, end - display_limit - warm_up)
        try:
            rows = self._read_parquet_rows(start, end - start)
        except Exception as error:
            raise SourceValidationError(
                f"{self.source_path.name}: cannot read Parquet page: {error}"
            ) from error
        bars = _parse_rows(
            self.source_path,
            [("time", "open", "high", "low", "close", "volume"), *rows],
            self.source_timezone,
            source_start_row=start + 2,
        )
        display_start = max(0, len(bars) - display_limit)
        return BarPage(bars, display_start, start + display_start > 0)

    def _read_parquet_rows(self, start: int, count: int) -> list[tuple[Any, ...]]:
        cache_path = self._parquet_cache_path
        assert cache_path is not None
        with closing(sqlite3.connect(cache_path)) as connection:
            return list(
                connection.execute(
                    "SELECT time, open, high, low, close, volume "
                    "FROM bars WHERE row_number >= ? AND row_number < ? ORDER BY row_number",
                    (start, start + count),
                ).fetchall()
            )

    def _xlsx_page(self, before: int | None, display_limit: int, warm_up: int) -> BarPage:
        try:
            workbook = load_workbook(self.source_path, read_only=True, data_only=True)
        except Exception as error:
            raise SourceValidationError(
                f"{self.source_path.name}: cannot read XLSX workbook: {error}"
            ) from error
        try:
            worksheet = workbook.active
            timestamps = self._timestamp_index()
            end = len(timestamps) if before is None else bisect_left(timestamps, before)
            start = max(0, end - display_limit - warm_up)
            header = tuple(cell.value for cell in worksheet[1])
            rows = [
                tuple(row)
                for row in worksheet.iter_rows(min_row=start + 2, max_row=end + 1, values_only=True)
            ]
            return self._page_from_rows([header, *rows], end, display_limit, warm_up)
        finally:
            workbook.close()

    def _page_from_rows(
        self,
        rows: list[tuple[Any, ...]],
        end: int,
        display_limit: int,
        warm_up: int,
    ) -> BarPage:
        source_start = end - (len(rows) - 1)
        bars = _parse_rows(
            self.source_path, rows, self.source_timezone, source_start_row=source_start + 2
        )
        display_start = max(0, len(bars) - display_limit)
        return BarPage(bars, display_start, source_start + display_start > 0)

    def _timestamp_index(self) -> tuple[int, ...]:
        if self._timestamps is not None:
            return self._timestamps
        suffix = self.source_path.suffix.lower()
        if suffix == ".csv":
            timestamps = self._build_csv_index()
        elif suffix == ".xlsx":
            timestamps = self._build_xlsx_index()
        elif suffix == ".parquet":
            timestamps = self._build_parquet_index()
        else:
            raise SourceValidationError(f"{self.source_path.name}: unsupported source format")
        _validate_timestamp_index(self.source_path, timestamps)
        self._timestamps = timestamps
        return timestamps

    def _build_xlsx_index(self) -> tuple[int, ...]:
        try:
            workbook = load_workbook(self.source_path, read_only=True, data_only=True)
        except Exception as error:
            raise SourceValidationError(
                f"{self.source_path.name}: cannot index XLSX workbook: {error}"
            ) from error
        try:
            worksheet = workbook.active
            columns = _columns(self.source_path, tuple(cell.value for cell in worksheet[1]))
            time_column = columns["time"] + 1
            return tuple(
                _timestamp(self.source_path, row_number, row[0], self.source_timezone)
                for row_number, row in enumerate(
                    worksheet.iter_rows(
                        min_row=2,
                        min_col=time_column,
                        max_col=time_column,
                        values_only=True,
                    ),
                    start=2,
                )
            )
        finally:
            workbook.close()

    def _build_parquet_index(self) -> tuple[int, ...]:
        try:
            cache_path = self._parquet_cache_path
            assert cache_path is not None
            with closing(sqlite3.connect(cache_path)) as connection:
                return tuple(
                    int(row[0])
                    for row in connection.execute("SELECT time FROM bars ORDER BY row_number")
                )
        except SourceValidationError:
            raise
        except Exception as error:
            raise SourceValidationError(
                f"{self.source_path.name}: cannot index Parquet file: {error}"
            ) from error

    def _canonical_parquet_columns(self, parquet_file: pq.ParquetFile) -> list[str]:
        names = parquet_file.schema_arrow.names
        _columns(self.source_path, tuple(names))
        volume_name = "volume" if "volume" in names else "tick_volume"
        return [*REQUIRED_PRICE_COLUMNS, volume_name]

    def _prepare_parquet_cache(self) -> Path:
        source_stat = self.source_path.stat()
        resolved_source = str(self.source_path.resolve())
        identity_values = (
            resolved_source,
            source_stat.st_mtime_ns,
            source_stat.st_size,
            str(self.source_timezone),
            PARQUET_CACHE_VERSION,
        )
        identity = repr(identity_values).encode()
        cache_directory = Path(tempfile.gettempdir()) / "breakout-research-parquet"
        cache_directory.mkdir(parents=True, exist_ok=True)
        cache_path = cache_directory / f"{hashlib.sha256(identity).hexdigest()}.sqlite3"
        self._clean_parquet_cache(cache_directory, resolved_source, cache_path)
        if self._valid_parquet_cache(cache_path, identity_values):
            return cache_path
        cache_path.unlink(missing_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f"{cache_path.stem}.", suffix=".tmp", dir=cache_directory
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        temporary_path.unlink()
        try:
            parquet_file = pq.ParquetFile(self.source_path)
            columns = self._canonical_parquet_columns(parquet_file)
            connection = sqlite3.connect(temporary_path)
            try:
                connection.execute(
                    "CREATE TABLE bars (row_number INTEGER PRIMARY KEY, time INTEGER NOT NULL, "
                    "open TEXT, high TEXT, low TEXT, close TEXT, volume TEXT)"
                )
                connection.execute(
                    "CREATE TABLE metadata (source_path TEXT, source_mtime_ns INTEGER, "
                    "source_size INTEGER, timezone TEXT, cache_version INTEGER)"
                )
                connection.execute("INSERT INTO metadata VALUES (?, ?, ?, ?, ?)", identity_values)
                connection.execute(f"PRAGMA user_version = {PARQUET_CACHE_VERSION}")
                row_number = 0
                for batch in parquet_file.iter_batches(columns=columns, batch_size=1_000):
                    cached_rows = []
                    for raw_row in zip(*(column.to_pylist() for column in batch.columns)):
                        timestamp = _timestamp(
                            self.source_path,
                            row_number + 2,
                            raw_row[0],
                            self.source_timezone,
                        )
                        cached_rows.append(
                            (
                                row_number,
                                timestamp,
                                *(_parquet_cell(value) for value in raw_row[1:]),
                            )
                        )
                        row_number += 1
                    connection.executemany(
                        "INSERT INTO bars VALUES (?, ?, ?, ?, ?, ?, ?)", cached_rows
                    )
                connection.commit()
            finally:
                connection.close()
            if self._valid_parquet_cache(cache_path, identity_values):
                temporary_path.unlink(missing_ok=True)
            else:
                try:
                    temporary_path.replace(cache_path)
                except OSError:
                    if not self._valid_parquet_cache(cache_path, identity_values):
                        raise
                    temporary_path.unlink(missing_ok=True)
        except Exception as error:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass
            if isinstance(error, SourceValidationError):
                raise
            raise SourceValidationError(
                f"{self.source_path.name}: cannot prepare Parquet page store: {error}"
            ) from error
        return cache_path

    @staticmethod
    def _valid_parquet_cache(path: Path, identity: tuple[object, ...]) -> bool:
        if not path.is_file():
            return False
        try:
            with closing(sqlite3.connect(f"file:{path}?mode=ro", uri=True)) as connection:
                if connection.execute("PRAGMA quick_check").fetchone() != ("ok",):
                    return False
                if connection.execute("PRAGMA user_version").fetchone() != (PARQUET_CACHE_VERSION,):
                    return False
                if tuple(row[1] for row in connection.execute("PRAGMA table_info(bars)")) != (
                    "row_number",
                    "time",
                    "open",
                    "high",
                    "low",
                    "close",
                    "volume",
                ):
                    return False
                if connection.execute("SELECT COUNT(*) FROM metadata").fetchone() != (1,):
                    return False
                if connection.execute("SELECT * FROM metadata").fetchone() != identity:
                    return False
                count, minimum, maximum = connection.execute(
                    "SELECT COUNT(*), MIN(row_number), MAX(row_number) FROM bars"
                ).fetchone()
                return (count == 0 and minimum is None and maximum is None) or (
                    minimum == 0 and maximum == count - 1
                )
        except (OSError, sqlite3.Error):
            return False

    @staticmethod
    def _clean_parquet_cache(cache_directory: Path, source_path: str, current_cache: Path) -> None:
        now = time.time()
        candidates = list(cache_directory.iterdir())[:PARQUET_CACHE_CLEANUP_LIMIT]
        for candidate in candidates:
            try:
                if candidate.suffix == ".tmp" and now - candidate.stat().st_mtime > 86_400:
                    candidate.unlink(missing_ok=True)
                    continue
                if candidate.suffix != ".sqlite3" or candidate == current_cache:
                    continue
                with closing(sqlite3.connect(f"file:{candidate}?mode=ro", uri=True)) as connection:
                    cached_source = connection.execute(
                        "SELECT source_path FROM metadata"
                    ).fetchone()
                if cached_source == (source_path,):
                    candidate.unlink(missing_ok=True)
            except (OSError, sqlite3.Error):
                continue


@dataclass(frozen=True, slots=True)
class DatasetSelection:
    symbol: str
    timeframe: str
    source_path: Path


class SourceCatalog:
    """A filename-only index; bar files are opened only by ``load_bars``."""

    def __init__(self, selections: Iterable[DatasetSelection]) -> None:
        self._selections = {(entry.symbol, entry.timeframe): entry for entry in selections}

    def symbols(self) -> list[dict[str, object]]:
        timeframes_by_symbol: dict[str, list[str]] = {}
        for symbol, timeframe in self._selections:
            timeframes_by_symbol.setdefault(symbol, []).append(timeframe)
        return [
            {"symbol": symbol, "timeframes": sorted(timeframes)}
            for symbol, timeframes in sorted(timeframes_by_symbol.items())
        ]

    def selection(self, symbol: str, timeframe: str) -> DatasetSelection:
        selection = self._selections.get((symbol, timeframe))
        if selection is None:
            raise SourceValidationError(f"{symbol}/{timeframe}: dataset is not available")
        return selection


def discover_source_catalog(data_root: Path) -> SourceCatalog:
    if not data_root.is_dir():
        raise SourceValidationError(f"{data_root}: data root does not exist or is not a directory")
    try:
        files = [path for path in data_root.rglob("*") if path.is_file()]
    except OSError as error:
        raise SourceValidationError(f"{data_root}: cannot enumerate data root: {error}") from error
    selections = [
        _selection_from_filename(path)
        for path in files
        if path.suffix.lower() in SUPPORTED_SOURCE_SUFFIXES
    ]
    catalog = SourceCatalog(selection for selection in selections if selection is not None)
    if not catalog.symbols():
        raise SourceValidationError(
            f"{data_root}: no CSV, Parquet, or XLSX symbol_timeframe files found"
        )
    if len(catalog._selections) != len(
        [selection for selection in selections if selection is not None]
    ):
        raise SourceValidationError(f"{data_root}: duplicate symbol/timeframe filenames found")
    return catalog


def _selection_from_filename(source_path: Path) -> DatasetSelection | None:
    stem = source_path.stem
    for suffix in EXPORT_NAME_SUFFIXES:
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break
    for separator in ("_", "-"):
        if separator in stem:
            symbol, timeframe = stem.rsplit(separator, 1)
            if symbol and timeframe:
                return DatasetSelection(symbol, timeframe, source_path)
    return None


def load_bars(
    source_path: Path, source_timezone: str = "UTC", dataset_id: str | None = None
) -> BarStore:
    timezone_info = _source_timezone(source_path, source_timezone)
    if not source_path.is_file():
        raise SourceValidationError(f"{source_path.name}: source file does not exist")
    if source_path.suffix.lower() == ".csv":
        _validate_csv_rows(source_path)
        frame = _read_tabular(source_path, "CSV", pd.read_csv)
        return _store_from_frame(source_path, frame, timezone_info, dataset_id)
    if source_path.suffix.lower() == ".parquet":
        frame = _read_tabular(source_path, "Parquet", pd.read_parquet)
        return _store_from_frame(source_path, frame, timezone_info, dataset_id)
    if source_path.suffix.lower() == ".xlsx":
        bars = _load_workbook(source_path, timezone_info)
        return BarStore(bars, dataset_id=dataset_id or _dataset_id(source_path))
    raise SourceValidationError(f"{source_path.name}: unsupported source format")


def _read_tabular(source_path: Path, format_name: str, reader: Any) -> pd.DataFrame:
    try:
        frame = reader(source_path)
    except Exception as error:
        raise SourceValidationError(
            f"{source_path.name}: cannot read {format_name} file: {error}"
        ) from error
    if not isinstance(frame, pd.DataFrame):
        raise SourceValidationError(f"{source_path.name}: {format_name} reader returned no table")
    return frame


def _validate_csv_rows(source_path: Path) -> None:
    try:
        with source_path.open(newline="", encoding="utf-8-sig") as source_file:
            rows = csv.reader(source_file)
            header = next(rows, [])
            if len(header) != len(set(header)):
                raise SourceValidationError(
                    f"{source_path.name}: duplicate column names are not allowed"
                )
            for row_number, row in enumerate(rows, start=2):
                if len(row) != len(header):
                    raise SourceValidationError(
                        f"{source_path.name}, CSV row {row_number}: expected "
                        f"{len(header)} fields, found {len(row)}"
                    )
    except SourceValidationError:
        raise
    except (OSError, UnicodeError, csv.Error) as error:
        raise SourceValidationError(f"{source_path.name}: cannot read CSV file: {error}") from error


def _store_from_frame(
    source_path: Path,
    frame: pd.DataFrame,
    source_timezone: ZoneInfo,
    dataset_id: str | None,
    source_start_row: int = 2,
) -> BarStore:
    if not frame.columns.is_unique:
        raise SourceValidationError(f"{source_path.name}: duplicate column names are not allowed")
    names = [str(column).strip() for column in frame.columns]
    if len(names) != len(set(names)):
        raise SourceValidationError(f"{source_path.name}: duplicate column names are not allowed")
    normalized_frame = frame.copy()
    normalized_frame.columns = pd.Index(names)
    rows = normalized_frame.itertuples(index=False, name=None)
    bars = _parse_rows(
        source_path,
        chain((tuple(names),), rows),
        source_timezone,
        source_start_row=source_start_row,
    )
    canonical_names = {*REQUIRED_PRICE_COLUMNS, "volume", "tick_volume"}
    imported_names = [column for column in names if column not in canonical_names]
    imported = normalized_frame.loc[:, imported_names].copy()
    imported.index = pd.Index(
        [
            _timestamp(source_path, row, value, source_timezone)
            for row, value in enumerate(normalized_frame["time"], start=source_start_row)
        ],
        dtype="int64",
        name="time",
    )
    return BarStore(bars, imported, dataset_id or _dataset_id(source_path))


def _dataset_id(source_path: Path) -> str:
    selection = _selection_from_filename(source_path)
    return (
        f"{selection.symbol}/{selection.timeframe}" if selection is not None else source_path.stem
    )


def _validated_imported_features(
    frame: pd.DataFrame | None, timestamps: tuple[int, ...], dataset_id: str
) -> pd.DataFrame:
    if frame is None:
        return pd.DataFrame(index=pd.Index(timestamps, dtype="int64", name="time"))
    if not frame.columns.is_unique:
        raise SourceValidationError(f"{dataset_id}: imported feature names must be unique")
    if tuple(frame.index) != timestamps:
        raise SourceValidationError(
            f"{dataset_id}: imported feature timestamps must align exactly with canonical bars"
        )
    if any(
        str(column) in REQUIRED_PRICE_COLUMNS or str(column) in {"volume", "tick_volume"}
        for column in frame.columns
    ):
        raise SourceValidationError(
            f"{dataset_id}: imported features must not overwrite canonical columns"
        )
    validated = frame.copy()
    validated.index = pd.Index(timestamps, dtype="int64", name="time")
    validated.attrs["dataset_id"] = dataset_id
    validated.attrs["feature_sources"] = {str(column): "imported" for column in validated.columns}
    return validated


def _load_workbook(source_path: Path, source_timezone: ZoneInfo) -> tuple[Bar, ...]:
    try:
        workbook = load_workbook(source_path, read_only=True, data_only=True)
    except Exception as error:
        raise SourceValidationError(
            f"{source_path.name}: cannot read XLSX workbook: {error}"
        ) from error
    try:
        try:
            worksheet = workbook.active
            bars = _parse_rows(source_path, worksheet.iter_rows(values_only=True), source_timezone)
        except SourceValidationError:
            raise
        except Exception as error:
            raise SourceValidationError(
                f"{source_path.name}: cannot read XLSX worksheet: {error}"
            ) from error
    finally:
        workbook.close()
    return bars


def _parse_rows(
    source_path: Path,
    rows: Iterable[Sequence[Any]],
    source_timezone: ZoneInfo,
    source_start_row: int = 2,
) -> tuple[Bar, ...]:
    row_iterator = iter(rows)
    header = next(row_iterator, None)
    columns = _columns(source_path, header)
    bars = tuple(
        _parse_bar(source_path, row_number, row, columns, source_timezone)
        for row_number, row in enumerate(row_iterator, start=source_start_row)
    )
    _validate_order(source_path, bars, source_start_row)
    return bars


def _source_timezone(source_path: Path, source_timezone: str) -> ZoneInfo:
    try:
        return ZoneInfo(source_timezone)
    except (ValueError, ZoneInfoNotFoundError) as error:
        raise SourceValidationError(
            f"{source_path.name}: source timezone {source_timezone!r} is not a valid IANA timezone"
        ) from error


def _columns(source_path: Path, header: Sequence[Any] | None) -> dict[str, int]:
    names = [str(value).strip() if value is not None else "" for value in (header or ())]
    missing = [name for name in REQUIRED_PRICE_COLUMNS if name not in names]
    if missing:
        raise SourceValidationError(
            f"{source_path.name}, spreadsheet row 1, column header: missing {missing}"
        )
    volume_name = (
        "volume" if "volume" in names else "tick_volume" if "tick_volume" in names else None
    )
    if volume_name is None:
        raise SourceValidationError(
            f"{source_path.name}, spreadsheet row 1, column header: missing volume or tick_volume"
        )
    return {name: names.index(name) for name in (*REQUIRED_PRICE_COLUMNS, volume_name)}


def _parse_bar(
    source_path: Path,
    row_number: int,
    row: Sequence[Any],
    columns: dict[str, int],
    source_timezone: ZoneInfo,
) -> Bar:
    timestamp = _timestamp(source_path, row_number, _cell(row, columns["time"]), source_timezone)
    values = {
        name: _number(source_path, row_number, name, _cell(row, column))
        for name, column in columns.items()
        if name != "time"
    }
    volume = values.pop("volume", values.pop("tick_volume", 0.0))
    if volume < 0:
        _error(source_path, row_number, "volume", volume, "must not be negative")
    if values["high"] < max(values["open"], values["close"]):
        _error(source_path, row_number, "high", values["high"], "must be at least open and close")
    if values["low"] > min(values["open"], values["close"]):
        _error(source_path, row_number, "low", values["low"], "must be at most open and close")
    if values["low"] > values["high"]:
        _error(source_path, row_number, "low", values["low"], "must not exceed high")
    return Bar(timestamp, values["open"], values["high"], values["low"], values["close"], volume)


def _cell(row: Sequence[Any], index: int) -> Any:
    return row[index] if index < len(row) else None


def _timestamp(source_path: Path, row: int, value: Any, source_timezone: ZoneInfo) -> int:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str) and value.lstrip("+-").isdigit():
        return _unix_timestamp(source_path, row, value, int(value))
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            _error(source_path, row, "time", value, "is not an ISO-8601 timestamp")
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        if not isfinite(float(value)) or int(value) != value:
            _error(
                source_path,
                row,
                "time",
                value,
                "must be an integer Unix seconds or milliseconds value",
            )
        return _unix_timestamp(source_path, row, value, int(value))
    else:
        _error(source_path, row, "time", value, "is not a timestamp")
    try:
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=source_timezone)
        utc_value = parsed.astimezone(UTC)
        if utc_value.microsecond:
            _error(source_path, row, "time", value, "loses sub-second precision")
        return int(utc_value.timestamp())
    except SourceValidationError:
        raise
    except (OverflowError, OSError, ValueError) as error:
        _error(source_path, row, "time", value, f"cannot convert to UTC Unix seconds: {error}")


def _unix_timestamp(source_path: Path, row: int, value: Any, raw: int) -> int:
    if abs(raw) >= 100_000_000_000 and raw % 1000:
        _error(source_path, row, "time", value, "loses sub-second precision")
    return raw // 1000 if abs(raw) >= 100_000_000_000 else raw


def _number(source_path: Path, row: int, column: str, value: Any) -> float:
    if value is None or isinstance(value, bool):
        _error(source_path, row, column, value, "must be a finite number")
    try:
        number = float(value)
    except (TypeError, ValueError):
        _error(source_path, row, column, value, "must be a finite number")
    if not isfinite(number):
        _error(source_path, row, column, value, "must be a finite number")
    return number


def _parquet_cell(value: Any) -> str | None:
    return None if value is None else str(value)


def _validate_order(source_path: Path, bars: tuple[Bar, ...], source_start_row: int = 2) -> None:
    for row_number, (previous, current) in enumerate(
        zip(bars, bars[1:]), start=source_start_row + 1
    ):
        if current.time == previous.time:
            _error(
                source_path, row_number, "time", current.time, "duplicates the previous timestamp"
            )
        if current.time < previous.time:
            _error(
                source_path,
                row_number,
                "time",
                current.time,
                "is earlier than the previous timestamp",
            )


def _validate_timestamp_index(source_path: Path, timestamps: tuple[int, ...]) -> None:
    for row_number, (previous, current) in enumerate(zip(timestamps, timestamps[1:]), start=3):
        if current == previous:
            _error(source_path, row_number, "time", current, "duplicates the previous timestamp")
        if current < previous:
            _error(
                source_path,
                row_number,
                "time",
                current,
                "is earlier than the previous timestamp",
            )


def _error(source_path: Path, row: int, column: str, value: Any, reason: str) -> NoReturn:
    raise SourceValidationError(
        f"{source_path.name}, spreadsheet row {row}, column {column}, value {value!r}: {reason}"
    )
